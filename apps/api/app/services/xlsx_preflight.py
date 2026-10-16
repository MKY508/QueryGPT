"""Read-only XLSX staging; no workbook-sized DataFrame or formula execution."""

from __future__ import annotations

import math
import re
from contextlib import ExitStack, closing
from datetime import date, datetime
from itertools import islice, zip_longest
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

import openpyxl
import pandas as pd

from app.core.config import settings
from app.services.data_preflight import PreflightInputError
from app.services.semantic_field_roles import has_identifier_semantics
from app.services.tabular_preflight import (
    _BATCH_BYTES,
    _BATCH_CELLS,
    _BATCH_ROWS,
    _MAX_COLUMNS,
    _NULL_VALUES,
    _RECORD_CHARACTERS,
    StagedTable,
    StringBatchWriter,
    _issue,
    run_staged_preflight,
)

_MAX_ARCHIVE_ENTRIES = 10_000
_MAX_SHEETS = 128
_MAX_ROWS = 1_048_576


def _check_archive(path: Path) -> None:
    """Guard allocations that openpyxl's read-only worksheet mode does not bound."""
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > _MAX_ARCHIVE_ENTRIES:
                raise ValueError("Excel 文件内容过多，请拆分工作表后再导入")
            names: set[str] = set()
            expanded = metadata = 0
            for entry in entries:
                if entry.filename in names:
                    raise ValueError("Excel 文件结构不完整，请重新保存后再导入")
                names.add(entry.filename)
                expanded += entry.file_size
                # Shared strings, styles, links, and other non-worksheet parts
                # may be loaded eagerly by the library, even in read-only mode.
                if not (
                    entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml")
                ):
                    metadata += entry.file_size
                if expanded > settings.FILE_PREPARATION_MAX_XLSX_EXPANDED_BYTES:
                    raise PreflightInputError(
                        "Excel 展开后过大，请拆分工作表或文件后再导入", code="xlsx_expanded_limit"
                    )
                if metadata > settings.FILE_PREPARATION_MAX_XLSX_METADATA_BYTES:
                    raise PreflightInputError(
                        "Excel 的文本或格式内容过多，请精简格式或另存为 CSV 后再导入",
                        code="xlsx_metadata_limit",
                    )
    except BadZipFile as exc:
        raise ValueError("无法读取这份 Excel，请重新保存为 XLSX 后再导入") from exc


def _rows(sheet):
    # Trust actual cells, not stale A1:A1 or inflated used-range metadata.
    sheet.reset_dimensions()
    iterator = sheet.iter_rows()
    try:
        for number, row in enumerate(iterator, 1):
            if number > _MAX_ROWS:
                raise ValueError("工作表行数超出 Excel 范围，请拆分文件后再导入")
            if len(row) > _MAX_COLUMNS and any(
                cell.value is not None for cell in row[_MAX_COLUMNS:]
            ):
                raise ValueError("工作表列数过多，请拆分文件后再导入")
            row = row[:_MAX_COLUMNS]
            if (
                sum(len(str(cell.value)) for cell in row if cell.value is not None)
                > _RECORD_CHARACTERS
            ):
                raise ValueError("工作表中有过长的记录，请缩短单元格内容后再导入")
            yield row
    finally:
        iterator.close()


def _probe(sheet) -> list[list[Any]]:
    with closing(_rows(sheet)) as rows:
        probe = [[cell.value for cell in row] for row in islice(rows, 12)]
    while probe and not any(value is not None for value in probe[-1]):
        probe.pop()
    return probe


def _text(value: Any, identifier: bool, number_format: str) -> str | None:
    if value is None or (isinstance(value, float) and not math.isfinite(value)):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if identifier and isinstance(value, (int, float)) and re.fullmatch(r"0{2,64}", number_format):
        if value >= 0 and int(value) == value:
            return str(int(value)).zfill(len(number_format))
    text = str(value)
    if not text or (not identifier and text in _NULL_VALUES):
        return None
    return text


def _stage_xlsx(
    source_path: Path, staged: Path, ordinal: str, steps: list[dict[str, Any]]
) -> StagedTable:
    from app.services import data_preflight as common

    _check_archive(source_path)
    issues: list[dict[str, Any]] = []
    operations: list[dict[str, Any]] = []
    drift: list[str] = []
    expected = {step["operation"] for step in steps}
    sheet_step = common._recipe_step(steps, "select_sheet")
    header_step = common._recipe_step(steps, "select_header")
    headerless = common._recipe_step(steps, "use_generated_headers") is not None

    with ExitStack() as stack:
        # Two lazy views let us distinguish absent formula caches from genuine
        # empty cells. Neither view calculates formulae or follows external links.
        cached = openpyxl.load_workbook(
            stack.enter_context(source_path.open("rb")),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        stack.callback(cached.close)
        if not cached.worksheets:
            raise ValueError("Excel 中没有可分析的工作表")
        if len(cached.worksheets) > _MAX_SHEETS:
            raise ValueError("Excel 工作表过多，请拆分文件后再导入")
        sheet_names = [sheet.title for sheet in cached.worksheets]
        # An explicit choice does not need to profile unrelated sheets. A wide
        # notes tab or a broken data row elsewhere must not veto this range.
        sheets_to_probe = (
            [cached[sheet_step["sheet"]]]
            if sheet_step and sheet_step["sheet"] in sheet_names
            else cached.worksheets
        )
        profiles, candidates = [], []
        for sheet in sheets_to_probe:
            probe = _probe(sheet)
            raw = pd.DataFrame(probe)
            header = (
                max(range(len(probe)), key=lambda index: common._header_score(raw.iloc[index]))
                if probe
                else 0
            )
            score = common._header_score(raw.iloc[header]) if probe else (0, 0)
            data_rows = sum(any(value is not None for value in row) for row in probe[header + 1 :])
            profiles.append(
                {
                    "name": sheet.title,
                    "header_row": header,
                    "textual_headers": score[0],
                    "unique_headers": score[1],
                    "sample_data_rows": data_rows,
                    "credible_table": score[1] >= 2 and data_rows >= 1,
                    "quality_score": score[1] * 2 + min(data_rows, 10),
                }
            )
            candidates.append((score, header, sheet.title))
        _, header, selected = max(candidates)
        if sheet_step:
            requested = sheet_step["sheet"]
            if requested in {item["name"] for item in profiles}:
                selected = requested
                header = next(item["header_row"] for item in profiles if item["name"] == selected)
            else:
                drift.append(f"原配方指定的工作表“{requested}”已不存在")
        probe = _probe(cached[selected])
        header_values = probe[header] if probe else []
        if header_step:
            requested = int(header_step["row"]) - 1
            # A recipe can name a physical row beyond the automatic 12-row
            # probe. Read only that row into memory, not every preceding row.
            with closing(_rows(cached[selected])) as rows:
                requested_cells = next(islice(rows, requested, requested + 1), ())
            requested_values = [cell.value for cell in requested_cells]
            if any(value is not None for value in requested_values):
                header = requested
                header_values = requested_values
            else:
                drift.append(f"原配方指定的第 {requested + 1} 行表头已不存在")
        if headerless:
            header = -1
        elif not header_values:
            raise ValueError(f"工作表“{selected}”中没有可分析的数据")

        formula_book = openpyxl.load_workbook(
            stack.enter_context(source_path.open("rb")),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        stack.callback(formula_book.close)
        formula_sheet = formula_book[selected]
        width = last_row = 0
        # Discover columns populated only late in the file without widening a
        # completed Arrow schema or discarding them. Trailing formatting is not data.
        with closing(_rows(formula_sheet)) as rows:
            for number, row in enumerate(rows, 1):
                if number <= header:
                    continue
                populated = [index + 1 for index, cell in enumerate(row) if cell.value is not None]
                if populated:
                    width = max(width, populated[-1])
                    last_row = number
        if not width or last_row <= header + 1:
            raise ValueError(f"工作表“{selected}”中没有可分析的数据")
        original_columns = (
            [None] * width if headerless else (header_values + [None] * width)[:width]
        )
        columns = common._dedupe_columns(original_columns)
        identifiers = [has_identifier_semantics(name) for name in columns]
        non_nulls, native_dates = [0] * width, [0] * width
        formula_count = 0
        with closing(_rows(cached[selected])) as values, closing(_rows(formula_sheet)) as formulae:
            with StringBatchWriter(staged, columns, ordinal) as writer:
                for number, (value_row, formula_row) in enumerate(zip_longest(values, formulae), 1):
                    if number > last_row:
                        break
                    if value_row is None or formula_row is None:
                        raise ValueError("工作表的内容发生变化，请重新导入")
                    if number <= header:
                        continue
                    # Include the header in cache/error validation, but never in data.
                    row: list[str | None] = []
                    for index in range(width):
                        cell = value_row[index] if index < len(value_row) else None
                        formula_cell = formula_row[index] if index < len(formula_row) else None
                        value = cell.value if cell else None
                        if cell and cell.data_type == "e":
                            raise PreflightInputError(
                                f"工作表“{selected}”中的 {cell.coordinate} 有错误值，请修正后再导入",
                                code="xlsx_cell_error",
                                sheet=selected,
                                cell=cell.coordinate,
                            )
                        if formula_cell and formula_cell.data_type == "f":
                            if value is None and (not cell or cell.data_type not in {"s", "str"}):
                                raise PreflightInputError(
                                    f"工作表“{selected}”中的 {formula_cell.coordinate} 没有已保存的公式结果，请先在 Excel 中计算并保存",
                                    code="xlsx_formula_result_missing",
                                    sheet=selected,
                                    cell=formula_cell.coordinate,
                                )
                            if number > header + 1:
                                formula_count += 1
                        if number == header + 1:
                            continue
                        text = _text(
                            value,
                            headerless or identifiers[index],
                            cell.number_format if cell else "General",
                        )
                        row.append(text)
                        if text is not None:
                            non_nulls[index] += 1
                            native_dates[index] += isinstance(value, (date, datetime))
                    if number > header + 1:
                        writer.append(row)

    def operation(name: str, **facts: Any) -> None:
        operations.append({"operation": name, **facts, "replayed": name in expected})

    operation("select_sheet", sheet=selected)
    if headerless:
        operation("use_generated_headers")
    elif header or "select_header" in expected:
        operation("select_header", row=header + 1)
        _issue(
            issues,
            "header_offset",
            f"已找到第 {header + 1} 行的真实表头",
            "表头之前的说明行未进入分析副本。",
            header,
            automatic=True,
        )
    if columns != original_columns or "normalize_column_names" in expected:
        operation("normalize_column_names", columns=columns)
    _issue(
        issues,
        "formula_results_review",
        f"有 {formula_count} 个公式结果需要核对",
        "读取的是文件中保存的结果，关键金额仍需核对是否已更新。",
        formula_count,
    )
    return StagedTable(
        rows=writer.count,
        columns=columns,
        operations=operations,
        issues=issues,
        drift=drift,
        literal_strings=headerless,
        native_dates=frozenset(
            name
            for index, name in enumerate(columns)
            if non_nulls[index] and non_nulls[index] == native_dates[index]
        ),
        reader={
            "engine": "xlsx_stream_v1",
            "selected_sheet": selected,
            "header_row": None if headerless else header + 1,
            "has_header": not headerless,
            "sheet_names": sheet_names,
            "sheet_profiles": profiles,
            "batch_rows": _BATCH_ROWS,
            "batch_cells": _BATCH_CELLS,
            "batch_bytes": _BATCH_BYTES,
            "formula_count": formula_count,
            "formula_results": "saved_values_only",
        },
    )


def run_xlsx_preflight(source_path: Path, output_dir: Path, steps: list[dict[str, Any]]):
    return run_staged_preflight(source_path, output_dir, steps, _stage_xlsx)
