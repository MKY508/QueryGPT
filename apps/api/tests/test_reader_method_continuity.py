"""Reader edits keep the same method, its history, and the last trusted result."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.api.v1 import projects
from app.core.config import settings
from app.db.tables import (
    PreflightReportRecord,
    ProjectDataSource,
    SanitationRecipeRecord,
    SanitationRecipeRevisionRecord,
)
from app.services.sanitation_contract import (
    SanitationContractError,
    replace_sanitation_reader_operations,
)
from app.services.sanitation_revisions import (
    append_sanitation_revision,
    ensure_sanitation_revision_head,
)


async def attach(client, db_session, tmp_path, monkeypatch, *, workbook=False):
    monkeypatch.setattr(settings, "WORKSPACE_ROOT", tmp_path / "workspace")
    project = (await client.post("/api/v1/projects", json={"name": "Reader edits"})).json()["data"]
    name, body = "orders.csv", b"order_id;amount\n001;10\n002;\n"
    if workbook:
        book = Workbook()
        book.active.title = "Orders"
        for row in [["order_id", "amount"], ["001", 10], ["002", None]]:
            book.active.append(row)
        later = book.create_sheet("Next month")
        for row in [["Monthly export"], ["order_id", "amount"], ["003", 30], ["004", None]]:
            later.append(row)
        book.create_sheet("Unrelated").append(["region", "sales"])
        book["Unrelated"].append(["West", 50])
        stream = BytesIO()
        book.save(stream)
        book.close()
        name, body = "orders.xlsx", stream.getvalue()
    response = await client.post(
        f"/api/v1/projects/{project['id']}/sources/files", files={"file": (name, body)}
    )
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    source = await db_session.get(ProjectDataSource, UUID(data["id"]))
    return f"/api/v1/projects/{project['id']}/sources/{data['id']}", source


async def reader_choice(client, url, *, sheet=None, **header):
    response = await client.get(
        url + ("/workbook" if sheet else "/csv"),
        params={"sheet": sheet} if sheet else {"delimiter": ";"},
    )
    assert response.status_code == 200, response.text
    preview = response.json()["data"]
    return {
        "source_revision": preview["source_revision"],
        "input_fingerprint": preview["input_fingerprint"],
        **(
            {"sheet": sheet}
            if sheet
            else {
                "kind": "csv",
                "delimiter": ";",
                "encoding": preview["encoding"],
            }
        ),
        **header,
    }


async def fill_amount(client, url):
    operations = [{"operation": "fill_missing", "column": "amount", "value": 0}]
    response = await client.post(url + "/cleaning/preview", json={"operations": operations})
    assert response.status_code == 200, response.text
    preview = response.json()["data"]
    applied = await client.post(
        url + "/cleaning/apply",
        json={
            "operations": operations,
            "expected_operations_hash": preview["operations_hash"],
            "expected_source_fingerprint": preview["source_fingerprint"],
            "expected_preview_output_fingerprint": preview["preview_output_fingerprint"],
            "expected_current_working_fingerprint": preview["current_working_fingerprint"],
            "expected_current_recipe_active_revision_id": preview[
                "current_recipe_active_revision_id"
            ],
        },
    )
    assert applied.status_code == 200, applied.text
    return applied.json()["data"]


async def current_method(db_session, source):
    recipe = (
        await db_session.execute(
            select(SanitationRecipeRecord).where(SanitationRecipeRecord.data_source_id == source.id)
        )
    ).scalar_one()
    await db_session.refresh(recipe)
    return recipe


def test_reader_replacement_preserves_transforms_and_validates_before_filtering():
    method = [
        {"operation": "read_csv", "delimiter": ",", "encoding": "utf-8-sig"},
        {"operation": "use_generated_headers"},
        {"operation": "fill_missing", "column": "amount", "value": 0},
    ]
    original = deepcopy(method)
    changed = replace_sanitation_reader_operations(
        method,
        [
            {"operation": "select_sheet", "sheet": "Orders"},
            {"operation": "select_header", "row": 14},
        ],
    )
    assert [step["operation"] for step in changed] == [
        "select_sheet",
        "select_header",
        "fill_missing",
    ]
    automatic = replace_sanitation_reader_operations(
        changed, [{"operation": "select_sheet", "sheet": "Orders"}]
    )
    assert [step["operation"] for step in automatic] == ["select_sheet", "fill_missing"]
    assert method == original
    with pytest.raises(SanitationContractError):
        replace_sanitation_reader_operations([{**method[0], "contract_version": 99}], [method[0]])


@pytest.mark.asyncio
async def test_csv_reader_edit_and_plain_prepare_keep_one_method_and_removable_manual_steps(
    client, db_session, tmp_path, monkeypatch
):
    url, source = await attach(client, db_session, tmp_path, monkeypatch)
    first = await client.post(
        url + "/preflight", json=await reader_choice(client, url, header_row=1)
    )
    assert first.status_code == 200, first.text
    await fill_amount(client, url)
    await db_session.refresh(source)
    manual_operations = deepcopy(source.profile_data["visual_cleaning"]["operations"])
    recipe = await current_method(db_session, source)
    recipe_id, old_head = recipe.id, recipe.active_revision_id
    old_operations = deepcopy(recipe.operations)
    original_file = Path(source.source_uri).read_bytes()

    async def never_guess(*args, **kwargs):
        pytest.fail("Re-preparing a source cannot guess a different source's method")

    monkeypatch.setattr(projects, "_matching_prior_source", never_guess)
    # Automatic header detection replaces the selected row, but not the fill.
    for payload in [await reader_choice(client, url), None]:
        response = await client.post(url + "/preflight", json=payload)
        assert response.status_code == 200, response.text
        await db_session.refresh(source)
        recipe = await current_method(db_session, source)
        assert recipe.id == recipe_id and source.profile_data["version"] == 1
        assert pq.read_table(source.working_uri).to_pylist()[-1] == {"order_id": "002", "amount": 0}
        assert source.profile_data["csv_options"]["delimiter"] == ";"
        assert source.profile_data["visual_cleaning"] == {
            "operations": manual_operations,
            "active_revision_id": str(recipe.active_revision_id),
        }
    previous = await db_session.get(SanitationRecipeRevisionRecord, old_head)
    assert previous.operations == old_operations and recipe.active_revision_id != old_head
    assert Path(source.source_uri).read_bytes() == original_file
    # Editor ownership follows the new head, so removing its fill is still possible.
    removed = await client.post(url + "/cleaning/preview", json={"operations": []})
    assert removed.status_code == 200, removed.text


@pytest.mark.asyncio
async def test_workbook_reader_edit_reuses_cleaning_and_rejects_incompatible_fields(
    client, db_session, tmp_path, monkeypatch
):
    url, source = await attach(client, db_session, tmp_path, monkeypatch, workbook=True)
    first = await client.post(
        url + "/preflight", json=await reader_choice(client, url, sheet="Orders", header_row=1)
    )
    assert first.status_code == 200, first.text
    await fill_amount(client, url)
    recipe = await current_method(db_session, source)
    recipe_id = recipe.id
    changed = await client.post(
        url + "/preflight", json=await reader_choice(client, url, sheet="Next month", header_row=2)
    )
    assert changed.status_code == 200, changed.text
    await db_session.refresh(source)
    recipe = await current_method(db_session, source)
    assert recipe.id == recipe_id
    assert source.profile_data["selected_sheet"] == "Next month"
    assert source.profile_data["header_row"] == 2
    assert pq.read_table(source.working_uri).to_pylist() == [
        {"order_id": "003", "amount": 30},
        {"order_id": "004", "amount": 0},
    ]
    trusted, head = source.working_uri, recipe.active_revision_id
    for choice in [
        await reader_choice(client, url, sheet="Unrelated", header_row=1),
        await reader_choice(client, url, sheet="Orders", headerless=True),
    ]:
        rejected = await client.post(url + "/preflight", json=choice)
        assert rejected.status_code == 409, rejected.text
        await db_session.refresh(source)
        recipe = await current_method(db_session, source)
        assert source.working_uri == trusted and recipe.active_revision_id == head


@pytest.mark.asyncio
async def test_reader_choices_bind_recipe_head_before_and_during_preparation(
    client, db_session, tmp_path, monkeypatch
):
    url, source = await attach(client, db_session, tmp_path, monkeypatch)
    assert (
        await client.post(url + "/preflight", json=await reader_choice(client, url, header_row=1))
    ).status_code == 200
    await db_session.refresh(source)
    recipe = await current_method(db_session, source)
    trusted = source.working_uri
    stale = await reader_choice(client, url, header_row=1)
    count = await db_session.scalar(select(func.count()).select_from(PreflightReportRecord))

    async def advance_head():
        active = await ensure_sanitation_revision_head(db_session, recipe)
        revision = await append_sanitation_revision(
            db_session,
            recipe,
            expected_active_revision_id=active.id,
            state="candidate",
            operations=active.operations,
            input_contract=active.input_contract,
            output_contract=active.output_contract,
            actor_source="user",
            reason="Restored method",
        )
        await db_session.commit()
        return revision.id

    head = await advance_head()
    rejected = await client.post(url + "/preflight", json=stale)
    assert rejected.status_code == 409, rejected.text
    assert await db_session.scalar(select(func.count()).select_from(PreflightReportRecord)) == count
    fresh = await reader_choice(client, url, header_row=1)
    assert fresh["source_revision"] != stale["source_revision"]
    runner = projects._run_preflight_in_thread

    async def change_method_while_preparing(*args, **kwargs):
        nonlocal head
        result = await runner(*args, **kwargs)
        head = await advance_head()
        return result

    monkeypatch.setattr(projects, "_run_preflight_in_thread", change_method_while_preparing)
    rejected = await client.post(url + "/preflight", json=fresh)
    assert rejected.status_code == 409, rejected.text
    await db_session.refresh(source)
    recipe = await current_method(db_session, source)
    assert source.working_uri == trusted and recipe.active_revision_id == head
    assert await db_session.scalar(select(func.count()).select_from(PreflightReportRecord)) == count


@pytest.mark.asyncio
async def test_reader_edit_keeps_replacement_pending_and_rechecks_schema(
    client, db_session, tmp_path, monkeypatch
):
    url, source = await attach(client, db_session, tmp_path, monkeypatch)
    assert (
        await client.post(url + "/preflight", json=await reader_choice(client, url, header_row=1))
    ).status_code == 200
    await db_session.refresh(source)
    old = ProjectDataSource(
        project_id=source.project_id,
        kind="file",
        name="prior.csv",
        format="csv",
        source_uri=source.source_uri,
        working_uri=source.working_uri,
        status="ready",
        fingerprint=source.fingerprint,
        profile_data={
            **source.profile_data,
            "schema": {
                "columns": [
                    {"name": "order_id", "type": "string"},
                    {"name": "missing_field", "type": "string"},
                ]
            },
        },
    )
    db_session.add(old)
    await db_session.flush()
    source.profile_data = {
        **source.profile_data,
        "is_current": False,
        "replacement_of": str(old.id),
        "activation_state": "pending_confirmation",
        "version": 2,
    }
    await db_session.commit()
    response = await client.post(
        url + "/preflight", json=await reader_choice(client, url, header_row=1)
    )
    assert response.status_code == 200, response.text
    await db_session.refresh(source)
    await db_session.refresh(old)
    assert source.profile_data["is_current"] is False and source.profile_data["version"] == 2
    assert source.profile_data["activation_state"] == "pending_confirmation"
    assert old.status == "ready" and old.profile_data["is_current"] is True
    snapshot = response.json()["data"]["source_snapshot"]
    assert snapshot["schema_drift"]["requires_confirmation"] is True
    assert (await current_method(db_session, source)).status == "needs_attention"
