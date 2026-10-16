"""Small generated workbooks; no real project files or installed-app launch."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4

import pandas as pd
import pyarrow.parquet as pq
import pytest
import xlrd
from openpyxl import Workbook
from sqlalchemy import func, select, update

from app.api.v1 import projects
from app.core.config import settings
from app.db.tables import PreflightReportRecord, ProjectDataSource
from app.services import data_preflight, workbook_preview
from app.services.data_preflight import fingerprint_file, run_preflight


def workbook():
    book = Workbook()
    orders = book.active
    orders.title = "Orders"
    orders.append(["order_id", "amount"])
    orders.append(["001", 10])
    orders.append(["002", 20])
    bad = book.create_sheet("Z Formula")
    bad.append(["region", "revenue", "tax", "total"])
    bad.append(["west", "=1+2", 3, 4])
    return book


def selection(preview, sheet="Orders"):
    return {
        "sheet": sheet,
        "input_fingerprint": preview["input_fingerprint"],
        "source_revision": preview["source_revision"],
    }


async def upload(client, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "WORKSPACE_ROOT", tmp_path / "projects")
    project = (await client.post("/api/v1/projects", json={"name": "Workbook navigation"})).json()[
        "data"
    ]
    body = BytesIO()
    book = workbook()
    # An unrelated sheet exceeds the preparation reader's column admission;
    # choosing Orders must not inspect these rows at all.
    book.create_sheet("Excluded notes").cell(1, 501, "Not selected")
    book.save(body)
    response = await client.post(
        f"/api/v1/projects/{project['id']}/sources/files",
        files={"file": ("orders.xlsx", body.getvalue())},
    )
    assert response.status_code == 200, response.text
    source = response.json()["data"]
    return f"/api/v1/projects/{project['id']}/sources/{source['id']}", source


def test_workbook_preview_is_bounded_lists_all_sheets_and_does_not_calculate_formulas(tmp_path):
    path = tmp_path / "preview.xlsx"
    book = workbook()
    for index in range(13):
        book.create_sheet(f"Extra {index}")
    book["Extra 12"].sheet_state = "hidden"
    sheet = book["Orders"]
    sheet["A2"] = 7
    sheet["A2"].number_format = "0000"
    sheet["B3"] = datetime(2026, 9, 5)
    sheet["C3"] = "x" * 400
    sheet["AO3"] = "omitted column"
    for index in range(15):
        sheet.append([index, index * 10])
    book.save(path)
    before = fingerprint_file(path)
    catalog = workbook_preview.preview_workbook(path)
    assert len(catalog["sheets"]) == 15 and catalog["rows"] == []
    assert catalog["sheets"][-1] == {"name": "Extra 12", "hidden": True}
    preview = workbook_preview.preview_workbook(path, "Orders")
    assert len(preview["rows"]) == 12 and max(map(len, preview["rows"])) == 40
    assert preview["rows"][1][0] == "0007"
    assert preview["rows"][2][1] == "2026-09-05T00:00:00"
    assert len(preview["rows"][2][2]) == 251
    assert preview["has_more_rows"] and preview["truncated_cells"] and preview["truncated_columns"]
    formula = workbook_preview.preview_workbook(path, "Z Formula")
    assert formula["missing_formula_results"] == 1 and formula["rows"][1][1] == "=1+2"
    assert fingerprint_file(path) == before
    with pytest.raises(ValueError, match="no longer available"):
        workbook_preview.preview_workbook(path, "Missing")


@pytest.mark.asyncio
async def test_worksheet_preview_precedes_preparation_and_explicit_choice_keeps_trusted_copy_on_failure(
    client, db_session, tmp_path, monkeypatch
):
    url, source_data = await upload(client, tmp_path, monkeypatch)
    source = await db_session.get(ProjectDataSource, UUID(source_data["id"]))
    original = fingerprint_file(Path(source.source_uri))
    preview = (await client.get(url + "/workbook", params={"sheet": "Orders"})).json()["data"]
    assert preview["selected_sheet"] == "Orders"
    assert await db_session.scalar(select(func.count()).select_from(PreflightReportRecord)) == 0
    assert source.working_uri is None and source.status == "attached"

    # Explicit range selection must not be replaced by a guessed prior recipe.
    async def unexpected_match(*args, **kwargs):
        pytest.fail("Manual worksheet choice cannot guess a prior source")

    monkeypatch.setattr(projects, "_matching_prior_source", unexpected_match)
    prepared = await client.post(url + "/preflight", json=selection(preview))
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()["data"]["source_snapshot"]["reader"]["selected_sheet"] == "Orders"
    await db_session.refresh(source)
    trusted = source.working_uri
    assert source.profile_data["selected_sheet"] == "Orders"
    assert pq.read_table(trusted).to_pylist() == [
        {"order_id": "001", "amount": 10},
        {"order_id": "002", "amount": 20},
    ]
    assert (await client.post(url + "/preflight", json=selection(preview))).status_code == 409
    bad_preview = (await client.get(url + "/workbook", params={"sheet": "Z Formula"})).json()[
        "data"
    ]
    rejected = await client.post(url + "/preflight", json=selection(bad_preview, "Z Formula"))
    assert rejected.status_code == 409
    await db_session.refresh(source)
    assert source.working_uri == trusted and source.profile_data["is_current"] is True
    assert fingerprint_file(Path(source.source_uri)) == original
    assert (
        await client.get(url.replace(source_data["project_id"], str(uuid4())) + "/workbook")
    ).status_code == 404


@pytest.mark.asyncio
async def test_workbook_selection_rejects_changed_file_and_an_intervening_source_write(
    client, db_session, tmp_path, monkeypatch
):
    url, source_data = await upload(client, tmp_path, monkeypatch)
    source = await db_session.get(ProjectDataSource, UUID(source_data["id"]))
    preview = (await client.get(url + "/workbook", params={"sheet": "Orders"})).json()["data"]
    changed = workbook()
    changed["Orders"]["B2"] = 99
    changed.save(source.source_uri)
    rejected = await client.post(url + "/preflight", json=selection(preview))
    assert (
        rejected.status_code == 409
        and rejected.json()["detail"]["code"] == "workbook_selection_changed"
    )
    assert await db_session.scalar(select(func.count()).select_from(PreflightReportRecord)) == 0
    fresh = (await client.get(url + "/workbook", params={"sheet": "Orders"})).json()["data"]
    run = projects._run_preflight_in_thread
    owned_other_copy = tmp_path / "other-writer.parquet"
    owned_other_copy.write_bytes(b"Other writer fixture")

    async def interleave(*args, **kwargs):
        result = await run(*args, **kwargs)
        await db_session.execute(
            update(ProjectDataSource)
            .where(ProjectDataSource.id == source.id)
            .values(working_uri=str(owned_other_copy), status="ready")
            .execution_options(synchronize_session=False)
        )
        await db_session.commit()
        return result

    monkeypatch.setattr(projects, "_run_preflight_in_thread", interleave)
    raced = await client.post(url + "/preflight", json=selection(fresh))
    assert (
        raced.status_code == 409 and raced.json()["detail"]["code"] == "workbook_selection_changed"
    )
    await db_session.refresh(source)
    assert source.working_uri == str(owned_other_copy)
    assert owned_other_copy.read_bytes() == b"Other writer fixture"


def test_legacy_preview_loads_only_the_chosen_sheet_and_releases_resources(tmp_path, monkeypatch):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"Legacy reader fixture")
    calls = []
    cells = [
        SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value="001"),
        SimpleNamespace(ctype=xlrd.XL_CELL_BOOLEAN, value=1),
        SimpleNamespace(ctype=xlrd.XL_CELL_ERROR, value=7),
    ]
    sheet = SimpleNamespace(visibility=1, nrows=1, row=lambda row: cells)
    book = SimpleNamespace(
        sheet_names=lambda: ["First", "Hidden"],
        sheet_by_name=lambda name: calls.append(name) or sheet,
        release_resources=lambda: calls.append("closed"),
        datemode=0,
    )
    monkeypatch.setattr(xlrd, "open_workbook", lambda *args, **kwargs: book)
    assert workbook_preview.preview_workbook(path)["rows"] == []
    assert calls == ["closed"]
    preview = workbook_preview.preview_workbook(path, "Hidden")
    assert preview["rows"] == [["001", "true", "#DIV/0!"]]
    assert calls == ["closed", "Hidden", "closed"]
    assert preview["sheets"][1]["hidden"] is True
    assert workbook_preview.preview_workbook(path, "Hidden", 2)["rows"] == []
    workbook_preview._READERS.acquire()
    workbook_preview._READERS.acquire()
    try:
        with pytest.raises(ValueError, match="busy"):
            workbook_preview.preview_workbook(path)
    finally:
        workbook_preview._READERS.release()
        workbook_preview._READERS.release()
    monkeypatch.setattr(settings, "FILE_UPLOAD_MAX_BYTES", 1)
    with pytest.raises(ValueError, match="too large"):
        workbook_preview.preview_workbook(path)
    monkeypatch.setattr(settings, "FILE_UPLOAD_MAX_BYTES", 1024)
    assert workbook_preview.preview_workbook(path)["rows"] == []


@pytest.mark.asyncio
async def test_explicit_late_header_is_published_and_unavailable_header_cannot_replace_it(
    client, db_session, tmp_path, monkeypatch
):
    url, source_data = await upload(client, tmp_path, monkeypatch)
    source = await db_session.get(ProjectDataSource, UUID(source_data["id"]))
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.append(["Report title", "Period", "A note", "Not column headings"])
    sheet.cell(2, 1, "=1+1")  # Excluded cover text must not need a formula cache.
    sheet.cell(15, 1, "order_id")
    sheet.cell(15, 2, "amount")
    sheet.append(["001", 10])
    sheet.append(["002", 20])
    book.save(source.source_uri)
    before = fingerprint_file(Path(source.source_uri))
    response = await client.get(url + "/workbook", params={"sheet": "Orders", "start_row": 13})
    assert response.status_code == 200, response.text
    sample = response.json()["data"]
    assert sample["start_row"] == 13 and sample["rows"][2] == ["order_id", "amount"]
    assert sample["nonempty_rows"] == [15, 16, 17]
    assert sample["missing_formula_results"] == 0
    body = {**selection(sample), "header_row": 15}
    prepared = await client.post(url + "/preflight", json=body)
    assert prepared.status_code == 200, prepared.text
    reader = prepared.json()["data"]["source_snapshot"]["reader"]
    assert reader["header_row"] == 15 and reader["selected_sheet"] == "Orders"
    await db_session.refresh(source)
    trusted = source.working_uri
    assert source.profile_data["header_row"] == 15
    assert pq.read_table(trusted).to_pylist() == [
        {"order_id": "001", "amount": 10},
        {"order_id": "002", "amount": 20},
    ]
    # Blank and out-of-range rows fail before another preparation is attempted.
    count = await db_session.scalar(select(func.count()).select_from(PreflightReportRecord))
    for row in [14, 999]:
        fresh = (
            await client.get(url + "/workbook", params={"sheet": "Orders", "start_row": row})
        ).json()["data"]
        rejected = await client.post(
            url + "/preflight", json={**selection(fresh), "header_row": row}
        )
        assert rejected.status_code == 409
        assert rejected.json()["detail"]["code"] == "workbook_header_unavailable"
    for row in [0, 1_000_001, True, "15"]:
        assert (
            await client.post(url + "/preflight", json={**body, "header_row": row})
        ).status_code == 422
    assert (
        await client.get(url + "/workbook", params={"sheet": "Orders", "start_row": 0})
    ).status_code == 422
    assert await db_session.scalar(select(func.count()).select_from(PreflightReportRecord)) == count
    await db_session.refresh(source)
    assert source.working_uri == trusted and fingerprint_file(Path(source.source_uri)) == before


def test_late_header_recipe_replay_works_after_blank_cover_rows_and_preview_pages_are_physical(
    tmp_path,
):
    path = tmp_path / "late.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.cell(25, 1, "order_id")
    sheet.cell(25, 2, "amount")
    for row in range(26, 45):
        sheet.cell(row, 1, f"00{row}")
        sheet.cell(row, 2, row)
    book.save(path)
    sample = workbook_preview.preview_workbook(path, "Orders", 20)
    assert sample["start_row"] == 20 and len(sample["rows"]) == 12
    assert sample["rows"][5] == ["order_id", "amount"] and sample["has_more_rows"]
    assert sample["nonempty_rows"] == list(range(25, 32))
    assert workbook_preview.preview_workbook(path, "Orders", 45)["rows"] == []
    assert workbook_preview.preview_workbook(path, "Orders", 1)["nonempty_rows"] == []
    steps = [
        {"operation": "select_sheet", "sheet": "Orders"},
        {"operation": "select_header", "row": 25},
    ]
    first = run_preflight(path, tmp_path / "first", recipe_operations=steps)
    replay = run_preflight(path, tmp_path / "replayed", recipe_operations=first.operations)
    assert first.source_snapshot["reader"]["header_row"] == 25
    assert pq.read_table(first.working_path).equals(pq.read_table(replay.working_path))
    assert pq.read_table(replay.working_path).num_rows == 19
    assert any(
        step["operation"] == "select_header" and step["row"] == 25 for step in replay.operations
    )
    with pytest.raises(ValueError, match="preview row"):
        workbook_preview.preview_workbook(path, "Orders", 0)


def test_legacy_header_beyond_probe_and_workbook_cleanup_on_reader_error(tmp_path, monkeypatch):
    calls = []

    class ExcelBook:
        sheet_names = ["Orders", "Unrelated"]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            calls.append("closed")

    frame = pd.DataFrame({"order_id": ["001"], "amount": [10]})
    fail = False

    def read_excel(book, **kwargs):
        assert kwargs["sheet_name"] == "Orders"
        calls.append(kwargs)
        if kwargs.get("skiprows") == 25:
            assert kwargs["keep_default_na"] is False
            return pd.DataFrame([["order_id", "amount"]])
        if kwargs.get("header") == 25:
            if fail:
                raise ValueError("Fixture read failure")
            return frame
        return pd.DataFrame([["Cover", None]])

    monkeypatch.setattr(pd, "ExcelFile", lambda path: ExcelBook())
    monkeypatch.setattr(pd, "read_excel", read_excel)
    steps = [
        {"operation": "select_sheet", "sheet": "Orders"},
        {"operation": "select_header", "row": 26},
    ]
    actual, header, reader = data_preflight._read_tabular(tmp_path / "legacy.xls", steps)
    assert actual is frame and header == 25 and reader["header_row"] == 26
    assert calls[-1] == "closed" and reader["recipe_replay_drift"] == []
    fail = True
    with pytest.raises(ValueError, match="Fixture read failure"):
        data_preflight._read_tabular(tmp_path / "legacy.xls", steps)
    assert calls[-1] == "closed"
