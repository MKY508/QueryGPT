"""Generated inputs for explicit headerless preparation; no real workbooks."""

from datetime import datetime
from io import BytesIO
from uuid import UUID

import pandas as pd
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from app.core.config import settings
from app.db.tables import ProjectDataSource
from app.services.data_preflight import _read_tabular, fingerprint_file, run_preflight
from app.services.sanitation_contract import (
    SanitationContractError,
    canonicalize_sanitation_operations,
)

GENERATED = [{"operation": "use_generated_headers"}]


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_headerless_keeps_first_record_literal_codes_and_replays_the_same_method(
    tmp_path, extension
):
    path = tmp_path / f"codes.{extension}"
    if extension == "csv":
        path.write_text("001,NA,10,true\n002,NULL,20,false\n", encoding="utf-8")
    else:
        book = Workbook()
        book.active.append(["001", "NA", 10, True, datetime(2026, 9, 5)])
        book.active.append(["002", "NULL", 20, False, datetime(2026, 9, 6)])
        book.save(path)
    before = fingerprint_file(path)
    first = run_preflight(path, tmp_path / "first", recipe_operations=GENERATED)
    rows = pq.read_table(first.working_path).to_pylist()
    assert len(rows) == 2
    assert [row["column_1"] for row in rows] == ["001", "002"]
    assert [row["column_2"] for row in rows] == ["NA", "NULL"]
    assert [row["column_3"] for row in rows] == [10, 20]
    assert [row["column_4"] for row in rows] == [True, False]
    if extension == "xlsx":
        assert rows[0]["column_5"] == datetime(2026, 9, 5)
    assert first.source_snapshot["original_rows"] == 2
    reader = first.source_snapshot["reader"]
    assert reader["has_header"] is False and reader["header_row"] is None
    assert any(step["operation"] == "use_generated_headers" for step in first.operations)
    assert not any(step["operation"] == "select_header" for step in first.operations)
    replay = run_preflight(path, tmp_path / "replay", recipe_operations=first.operations)
    assert pq.read_table(first.working_path).equals(pq.read_table(replay.working_path))
    assert fingerprint_file(path) == before


def test_headerless_single_row_is_data_and_first_row_formula_errors_remain_guarded(tmp_path):
    path = tmp_path / "single.xlsx"
    book = Workbook()
    book.active.append([7, 99])
    book.active["A1"].number_format = "0000"
    book.save(path)
    output = tmp_path / "prepared"
    result = run_preflight(path, output, recipe_operations=GENERATED)
    assert pq.read_table(result.working_path).to_pylist() == [{"column_1": "0007", "column_2": 99}]
    trusted = result.working_path.read_bytes()
    book.active["B1"] = "=1+2"
    book.save(path)
    with pytest.raises(ValueError, match="公式结果"):
        run_preflight(path, output, recipe_operations=GENERATED)
    assert result.working_path.read_bytes() == trusted


def test_headerless_contract_conflicts_and_incompatible_inputs_fail_before_reading(tmp_path):
    canonical = canonicalize_sanitation_operations(GENERATED)
    assert canonical == [{"operation": "use_generated_headers", "contract_version": 1}]
    with pytest.raises(SanitationContractError, match="不能同时"):
        canonicalize_sanitation_operations([*GENERATED, {"operation": "select_header", "row": 1}])
    for bad in [
        {"operation": "use_generated_headers", "row": 1},
        {"operation": "use_generated_headers", "contract_version": 2},
    ]:
        with pytest.raises(SanitationContractError):
            canonicalize_sanitation_operations([bad])
    for extension in ["json", "parquet"]:
        with pytest.raises(ValueError, match="仅适用于"):
            run_preflight(
                tmp_path / f"missing.{extension}", tmp_path / "absent", recipe_operations=GENERATED
            )
    assert not (tmp_path / "absent").exists()


def test_headerless_legacy_reader_preserves_stored_values_and_emits_generated_columns(
    tmp_path, monkeypatch
):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"Legacy reader facade")
    closed = []

    class ExcelBook:
        sheet_names = ["Orders"]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            closed.append(True)

    def read_excel(book, **kwargs):
        if "nrows" not in kwargs:
            assert kwargs["header"] is None and kwargs["dtype"] is object
            assert kwargs["keep_default_na"] is False
        return pd.DataFrame([["001", "NA", 10], ["002", "NULL", 20]], dtype=object)

    monkeypatch.setattr(pd, "ExcelFile", lambda path: ExcelBook())
    monkeypatch.setattr(pd, "read_excel", read_excel)
    frame, _, _ = _read_tabular(path, GENERATED)
    assert pd.api.types.is_integer_dtype(frame["column_3"].dtype)
    assert frame["column_1"].tolist() == ["001", "002"]
    result = run_preflight(path, tmp_path / "prepared", recipe_operations=GENERATED)
    assert closed == [True, True]
    assert pq.read_table(result.working_path).to_pylist() == [
        {"column_1": "001", "column_2": "NA", "column_3": 10},
        {"column_1": "002", "column_2": "NULL", "column_3": 20},
    ]
    assert result.source_snapshot["reader"]["has_header"] is False
    assert any(step["operation"] == "use_generated_headers" for step in result.operations)


@pytest.mark.asyncio
async def test_headerless_api_publishes_choice_and_rejects_conflicting_or_stale_selection(
    client, db_session, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "WORKSPACE_ROOT", tmp_path / "projects")
    project = (await client.post("/api/v1/projects", json={"name": "Headerless"})).json()["data"]
    book = Workbook()
    book.active.title = "Orders"
    book.active.append(["001", 10])
    book.active.append(["002", 20])
    content = BytesIO()
    book.save(content)
    uploaded = await client.post(
        f"/api/v1/projects/{project['id']}/sources/files",
        files={"file": ("orders.xlsx", content.getvalue())},
    )
    assert uploaded.status_code == 200, uploaded.text
    source_data = uploaded.json()["data"]
    url = f"/api/v1/projects/{project['id']}/sources/{source_data['id']}"
    preview = (await client.get(url + "/workbook", params={"sheet": "Orders"})).json()["data"]
    selection = {
        "sheet": "Orders",
        "headerless": True,
        "source_revision": preview["source_revision"],
        "input_fingerprint": preview["input_fingerprint"],
    }
    for extra in [{"header_row": 1}, {"headerless": "true"}]:
        assert (
            await client.post(url + "/preflight", json={**selection, **extra})
        ).status_code == 422
    result = await client.post(url + "/preflight", json=selection)
    assert result.status_code == 200, result.text
    assert result.json()["data"]["source_snapshot"]["reader"]["has_header"] is False
    source = await db_session.get(ProjectDataSource, UUID(source_data["id"]))
    await db_session.refresh(source)
    assert source.profile_data["has_header"] is False and source.profile_data["header_row"] is None
    assert pq.read_table(source.working_uri).num_rows == 2
    trusted = source.working_uri
    assert (await client.post(url + "/preflight", json=selection)).status_code == 409
    await db_session.refresh(source)
    assert source.working_uri == trusted
