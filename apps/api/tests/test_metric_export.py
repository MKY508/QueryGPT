from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.db.tables import Project
from app.services.report_metric_export import metric_export_rows


@pytest.mark.asyncio
async def test_metric_export_includes_exact_values_and_display_choices_without_formulas(
    client, db_session
):
    project = Project(name="指标导出")
    db_session.add(project)
    await db_session.commit()
    base = f"/api/v1/projects/{project.id}/reports"
    created = await client.post(
        base,
        json={
            "title": "经营指标",
            "pages": [
                {
                    "title": "概览",
                    "blocks": [
                        {
                            "block_type": "metric",
                            "title": "金额",
                            "source_kind": "manual",
                            "order_index": 0,
                            "content": {
                                "value": "9007199254740993.015",
                                "context": "=1+1",
                                "rows": [{"private": "secret"}],
                            },
                            "config": {
                                "metric_display": {
                                    "style": "currency",
                                    "currency": "USD",
                                    "decimals": 2,
                                    "suffix": "@SUM(A1)",
                                    "sql": "secret",
                                }
                            },
                        },
                        {
                            "block_type": "metric",
                            "title": "占比",
                            "source_kind": "manual",
                            "order_index": 1,
                            "content": {"value": 0.125},
                            "config": {"metric_display": {"style": "percent", "decimals": 2}},
                        },
                        {
                            "block_type": "metric",
                            "title": "空值",
                            "source_kind": "manual",
                            "order_index": 2,
                            "content": {"value": None, "amount": 999},
                        },
                    ],
                }
            ],
        },
    )
    assert created.status_code == 201, created.text
    report_id = created.json()["data"]["id"]
    exported = await client.get(base + f"/{report_id}/export.xlsx")
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content))
    assert book.sheetnames == ["报表概览", "金额", "占比", "空值"]
    assert [cell.value for cell in book["金额"][2]] == [
        "9007199254740993.015",
        "货币",
        2,
        "USD",
        "是",
        "'@SUM(A1)",
        "'=1+1",
    ]
    assert book["金额"]["A2"].data_type == "s"
    assert book["占比"]["A2"].value == 0.125  # Never multiply stored data by 100.
    assert book["空值"]["A2"].value is None
    assert all(
        cell.data_type != "f" and "secret" not in str(cell.value)
        for sheet in book
        for row in sheet
        for cell in row
    )
    document = (await client.get(base + f"/{report_id}")).json()["data"]
    assert (
        document["version"] == 1
        and document["pages"][0]["blocks"][0]["content"]["value"] == "9007199254740993.015"
    )


def test_metric_export_normalizes_unsupported_presentation_without_exposing_config():
    columns, rows = metric_export_rows(
        {"amount": "128万元"},
        {
            "metric_display": {
                "style": [],
                "decimals": True,
                "currency": {},
                "suffix": ["secret"],
            }
        },
    )
    assert len(columns) == 7 and rows == [["128万元", "原样", "", "", "", "", None]]
