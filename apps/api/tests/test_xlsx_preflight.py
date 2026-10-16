"""Small disposable workbooks covering the streamed Excel reader's boundaries."""

import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
import pyarrow.parquet as pq
import pytest
from httpx import AsyncClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.db.tables import PreflightReportRecord, ProjectDataSource
from app.services import tabular_preflight, xlsx_preflight
from app.services.data_preflight import fingerprint_file, run_preflight


def _rewrite_sheet(path: Path, transform) -> None:
    # Only generated test fixtures are rewritten; production readers are read-only.
    with ZipFile(BytesIO(path.read_bytes())) as source:
        with ZipFile(path, "w", compression=ZIP_DEFLATED) as target:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    content = transform(content.decode("utf-8")).encode("utf-8")
                target.writestr(entry.filename, content)


@pytest.mark.parametrize("dimension", ["A1:A1", "A1:XFD1048576"])
def test_streams_actual_cells_with_late_columns_and_preserves_excel_types(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    dimension: str,
):
    source = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_id", "happened", "active", "amount"])
    for index in range(20):
        sheet.append(
            [
                7 if index == 0 else "NA" if index == 1 else f"00{index + 100}",
                datetime(2026, 9, 1),
                index != 1,
                index + 10,
            ]
        )
    sheet["A2"].number_format = "0000"
    sheet["E20"] = "late column"
    sheet.append([7, datetime(2026, 9, 1), True, 10])
    sheet["A22"].number_format = "0000"
    workbook.save(source)
    _rewrite_sheet(
        source, lambda xml: re.sub(r'<dimension ref="[^"]+"', f'<dimension ref="{dimension}"', xml)
    )
    before = fingerprint_file(source)
    monkeypatch.setattr(
        pd, "read_excel", lambda *args, **kwargs: pytest.fail("No full Excel frame")
    )
    monkeypatch.setattr(pd, "ExcelFile", lambda *args, **kwargs: pytest.fail("Use a lazy workbook"))
    monkeypatch.setattr(tabular_preflight, "_BATCH_ROWS", 2)
    batches = []
    original_writer = pq.ParquetWriter

    class RecordingWriter(original_writer):
        def write_batch(self, batch, *args, **kwargs):
            batches.append(batch.num_rows)
            return super().write_batch(batch, *args, **kwargs)

    monkeypatch.setattr(pq, "ParquetWriter", RecordingWriter)
    result = run_preflight(source, tmp_path / "prepared")
    rows = pq.read_table(result.working_path).to_pylist()
    assert len(rows) == 20
    assert rows[0]["order_id"] == "0007"
    assert rows[1]["order_id"] == "NA"
    assert rows[0]["happened"] == datetime(2026, 9, 1)
    assert rows[0]["active"] is True and rows[1]["active"] is False
    assert rows[18]["column_5"] == "late column"
    assert sum(batches) == 21 and max(batches) <= 2
    assert result.source_snapshot["original_rows"] == 21
    assert result.source_snapshot["original_columns"] == 5
    assert result.source_snapshot["reader"]["engine"] == "xlsx_stream_v1"
    assert next(issue for issue in result.issues if issue["code"] == "duplicate_rows")["count"] == 1
    assert fingerprint_file(source) == before
    assert list((tmp_path / "prepared").iterdir()) == [result.working_path]


@pytest.mark.parametrize(
    "setting",
    ["FILE_PREPARATION_MAX_XLSX_METADATA_BYTES", "FILE_PREPARATION_MAX_XLSX_EXPANDED_BYTES"],
)
def test_rejects_archive_limits_before_loading_and_keeps_the_previous_copy(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    setting: str,
):
    source = tmp_path / "limited.xlsx"
    workbook = Workbook()
    workbook.active.append(["order_id", "amount"])
    workbook.active.append(["001", 10])
    workbook.save(source)
    output = tmp_path / "prepared"
    output.mkdir()
    previous = output / "analysis-ready.parquet"
    previous.write_bytes(b"trusted previous copy")
    monkeypatch.setattr(settings, setting, 1)
    monkeypatch.setattr(
        xlsx_preflight.openpyxl,
        "load_workbook",
        lambda *args, **kwargs: pytest.fail("Reject before loading"),
    )
    with pytest.raises(ValueError, match="Excel"):
        run_preflight(source, output)
    assert previous.read_bytes() == b"trusted previous copy"
    assert list(output.iterdir()) == [previous]


@pytest.mark.parametrize("value, message", [("=1+2", "公式结果"), ("#DIV/0!", "错误值")])
def test_late_uncached_formula_or_error_fails_without_publishing(
    tmp_path: Path,
    value: str,
    message: str,
):
    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_id", "amount"])
    for index in range(20):
        sheet.append([f"{index:04d}", index])
    sheet["B21"] = value
    workbook.save(source)
    output = tmp_path / "prepared"
    output.mkdir()
    previous = output / "analysis-ready.parquet"
    previous.write_bytes(b"trusted previous copy")
    with pytest.raises(ValueError, match=message):
        run_preflight(source, output)
    assert previous.read_bytes() == b"trusted previous copy"
    assert list(output.iterdir()) == [previous]


@pytest.mark.parametrize("cache, expected", [("<v>3</v>", 3), ("<v></v>", None)])
def test_reads_saved_formula_results_including_explicit_empty_strings(
    tmp_path: Path,
    cache: str,
    expected,
):
    source = tmp_path / "cached.xlsx"
    workbook = Workbook()
    workbook.active.append(["order_id", "amount"])
    workbook.active.append(["001", "=1+2"])
    workbook.active.append(["002", 10])
    workbook.save(source)

    def add_cache(xml):
        def replace(match):
            cell = re.sub(r"<v(?:\s*/>|>.*?</v>)", cache, match.group())
            if expected is None:
                cell = cell.replace('r="B2"', 'r="B2" t="str"')
            return cell

        return re.sub(r'<c r="B2".*?</c>', replace, xml)

    _rewrite_sheet(source, add_cache)
    result = run_preflight(source, tmp_path / "prepared")
    rows = pq.read_table(result.working_path).to_pylist()
    assert rows[0]["amount"] == expected
    assert result.source_snapshot["reader"]["formula_count"] == 1
    assert any(issue["code"] == "formula_results_review" for issue in result.issues)


@pytest.mark.asyncio
async def test_api_persists_actionable_formula_location_without_activating_source(
    client: AsyncClient,
    db_session: AsyncSession,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(settings, "WORKSPACE_ROOT", tmp_path / "projects")
    created = await client.post("/api/v1/projects", json={"name": "Formula review"})
    project_id = created.json()["data"]["id"]
    workbook = Workbook()
    workbook.active.title = "Orders"
    workbook.active.append(["order_id", "amount"])
    workbook.active.append(["001", "=1+2"])
    body = BytesIO()
    workbook.save(body)
    uploaded = await client.post(
        f"/api/v1/projects/{project_id}/sources/files",
        files={
            "file": (
                "formula.xlsx",
                body.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 200, uploaded.text
    source_id = uploaded.json()["data"]["id"]
    response = await client.post(f"/api/v1/projects/{project_id}/sources/{source_id}/preflight")
    assert response.status_code == 409
    report = (
        await db_session.execute(
            select(PreflightReportRecord).where(
                PreflightReportRecord.data_source_id == UUID(source_id)
            )
        )
    ).scalar_one()
    assert report.issues[0]["code"] == "preflight_failed"
    assert report.issues[0]["presentation_code"] == "xlsx_formula_result_missing"
    assert report.issues[0]["presentation_facts"] == {"sheet": "Orders", "cell": "B2"}
    assert "1+2" not in str(report.issues)
    source = await db_session.get(ProjectDataSource, UUID(source_id))
    assert source.status == "error" and source.working_uri is None
    assert source.profile_data["is_current"] is False
